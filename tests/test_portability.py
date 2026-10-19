from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from primer_build.workbooks import build_project_workbooks, SUMMARY_COLUMNS
from primer_build.project_pipeline import run_project, _design_one
from primer_build.codon_opt import optimize_codons
from primer_build.io import load_rules, read_structure_design
from primer_build.backbone import reconstruct_backbone
from Bio.Seq import Seq


class PortabilityTests(unittest.TestCase):
    def test_cys_labeling_template_matches_default_rules(self):
        rules = load_rules(ROOT / 'config/design_rules.json')
        expression = rules['expression']
        self.assertEqual(str(Seq(expression['fixed_c_terminal_dna']).translate()),
                         'GGSGGSGCSLEHHHHHH*')
        backbone = reconstruct_backbone(
            [ROOT / 'templates' / expression['backbone_template']],
            expression['fixed_c_terminal_dna'],
        )
        self.assertEqual(backbone.poi_insertion_index, 5200)
        self.assertEqual(
            backbone.sequence_without_poi[5200:5254],
            expression['fixed_c_terminal_dna'],
        )


    def test_workbook_roundtrip_layout_types_and_deduplication(self):
        with tempfile.TemporaryDirectory() as folder:
            primer = Path(folder) / 'primers.xlsx'
            summary = Path(folder) / 'summary.xlsx'
            row = {key: '' for key, _, _ in SUMMARY_COLUMNS}
            row.update(map='example.dna', proteinLength=80, initiationOpen=0.3833)
            build_project_workbooks(
                [{'name': '=literal', 'sequence': 'ac gt'},
                 {'name': 'alias', 'sequence': 'ACGT'},
                 {'name': 'second', 'sequence': 'TTAA'}], [row], primer, summary)
            wb = load_workbook(primer)
            self.assertEqual(wb.sheetnames, ['Primers'])
            sheet = wb.active
            self.assertEqual(list(sheet.values), [('Primer名称', "序列(5' to3')"), ('=literal', 'ACGT'), ('second', 'TTAA')])
            self.assertEqual(sheet['A2'].data_type, 's')
            self.assertEqual(sheet['B2'].number_format, '@')
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual(sheet.column_dimensions['B'].width, 108)
            self.assertEqual(sheet['A1'].fill.fgColor.rgb, '00D9E4F2')
            wb.close()
            wb = load_workbook(summary)
            self.assertEqual(wb.sheetnames, ['Summary'])
            self.assertEqual(tuple(wb.active.values)[0], tuple(c[1] for c in SUMMARY_COLUMNS))
            self.assertEqual(wb.active['D2'].value, 80)
            self.assertEqual(wb.active['P2'].value, 0.3833)
            self.assertEqual(wb.active['P2'].data_type, 'n')
            wb.close()

    def test_golden_reference_rejects_overwrite(self):
        for overwrite in (False, True):
            with self.assertRaisesRegex(ValueError, 'read-only golden'):
                run_project(ROOT, ROOT / 'ADN round1', 'A', 'ADN-1', 'A1-1-1-FW', overwrite=overwrite)

    def test_arbitrary_nonstandard_residues_encode_tag(self):
        rules = load_rules(ROOT / 'config/design_rules.json')
        for candidate in optimize_codons('MAXXG', rules['codon_optimization']):
            self.assertEqual(candidate.cds[6:12], 'TAGTAG')

    def test_mmcif_extension_and_chain_selection(self):
        content = '''data_chains
loop_
_entity_poly_seq.entity_id
_entity_poly_seq.num
_entity_poly_seq.mon_id
1 1 ALA
2 1 MET
2 2 SEP
loop_
_entity_poly.entity_id
_entity_poly.pdbx_strand_id
1 A
2 B
'''
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'two.mmcif'
            path.write_text(content)
            self.assertEqual(read_structure_design(path, 'test', protein_chain='B').expression_sequence, 'MX')
            with self.assertRaisesRegex(ValueError, 'not found'):
                read_structure_design(path, 'test', protein_chain='Z')


if __name__ == '__main__':
    unittest.main()
