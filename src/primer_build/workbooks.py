"""Portable, local XLSX export. No external runtime or preview dependency."""
from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SUMMARY_COLUMNS = (
    ('map', 'Map', 22), ('structure', 'Structure', 34), ('chain', 'Chain', 9),
    ('proteinLength', 'Protein (aa)', 12), ('cdsLength', 'CDS (bp)', 11),
    ('nonstandardResidues', 'Nonstandard residues', 14), ('fragments', 'Fragments', 10),
    ('designPrimers', 'Design primers', 14), ('longTierPrimers', 'Long-tier primers', 15),
    ('oeLengths', 'OE overlaps (bp)', 20), ('oeTm', 'OE Tm (°C)', 22),
    ('gibsonLengths', 'Gibson overlaps (bp)', 22), ('gibsonTm', 'Gibson Tm (°C)', 22),
    ('candidate', 'Candidate', 11), ('fivePrimeMfe', "5' MFE (kcal/mol)", 18),
    ('initiationOpen', 'RBS/ATG unpaired probability', 26),
    ('localMfeMin', 'Min local MFE (kcal/mol)', 24),
    ('localMfeStart', 'Min local MFE start (nt)', 24), ('sdLikeSites', 'SD-like sites', 24),
    ('expressionPenalty', 'Expression penalty', 18), ('expressionFlags', 'Expression flags', 28),
    ('status', 'Status', 10), ('notes', 'Notes', 42),
)


def _workbook(title, headers, rows, widths):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    primer = title == 'Primers'
    border = Border(**{key: Side(style='thin', color='C7C7C7' if primer else 'D9D9D9')
                       for key in ('left', 'right', 'top', 'bottom')})
    header_border = Border(**{key: Side(style='thin', color='000000' if primer else '7F7F7F')
                              for key in ('left', 'right', 'top', 'bottom')})
    for row in sheet:
        for cell in row:
            cell.font = Font(name='Calibri', size=11, color='000000')
            cell.border = border
            cell.alignment = Alignment(vertical='center', wrap_text=not primer and cell.column in (2, 19, 20, 21, 23))
            # Names and sequences are literal text, even if a name begins with '='.
            if isinstance(cell.value, str):
                cell.data_type = 's'
                cell.number_format = '@'
        sheet.row_dimensions[row[0].row].height = 21 if primer else 34
    for cell in sheet[1]:
        cell.font = Font(name='Calibri', bold=True, size=14 if primer else 11, color='000000')
        cell.fill = PatternFill('solid', fgColor='D9E4F2')
        cell.border = header_border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=not primer)
    sheet.row_dimensions[1].height = 46 if primer else 36
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return workbook


def _save_atomic(workbook: Workbook, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f'.{output.stem}-', suffix='.xlsx', dir=output.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
        workbook.close()


def build_project_workbooks(
    primer_rows: list[dict[str, str]],
    summary_rows: list[dict[str, object]],
    primer_output: Path,
    summary_output: Path,
) -> None:
    seen = set()
    primers = []
    for row in primer_rows:
        sequence = ''.join(row['sequence'].upper().split())
        if sequence not in seen:
            seen.add(sequence)
            primers.append((row['name'], sequence))
    _save_atomic(_workbook('Primers', ['Primer名称', "序列(5' to3')"], primers, [24, 108]), primer_output)
    _save_atomic(_workbook('Summary', [c[1] for c in SUMMARY_COLUMNS],
                          [[row[c[0]] for c in SUMMARY_COLUMNS] for row in summary_rows],
                          [c[2] for c in SUMMARY_COLUMNS]), summary_output)
