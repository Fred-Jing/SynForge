"""Independently read back one generated batch and its source provenance."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import defaultdict
from pathlib import Path

from Bio import SeqIO
from Bio.Seq import Seq
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from primer_build.io import get_poi_feature, load_rules, read_structure_design  # noqa: E402
from primer_build.project_pipeline import parse_numbering  # noqa: E402


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def circular_count(sequence: str, motif: str) -> int:
    doubled = sequence + sequence[: len(motif) - 1]
    return sum(doubled.startswith(motif, index) for index in range(len(sequence)))


def validate(project: Path) -> dict:
    manifest = json.loads((project / "run_manifest.json").read_text())
    provenance = json.loads((project / "input_provenance.json").read_text())
    rules = load_rules(Path(manifest["rules_file"]))
    plasmid_prefix = rules["expression"].get("plasmid_name_prefix", "PET22B")
    tail_dna = rules["expression"]["fixed_c_terminal_dna"]
    expected_tail = rules["expression"]["fixed_c_terminal_translation"]
    numbering = parse_numbering(
        manifest["parameters"]["map_start"], manifest["parameters"]["primer_start"]
    )
    for relative, expected_hash in (
        manifest["input_sha256"] | manifest["output_sha256"]
    ).items():
        if digest(project / relative) != expected_hash:
            raise ValueError(f"Manifest SHA-256 mismatch: {relative}")
    if digest(Path(manifest["rules_file"])) != manifest["rules_sha256"]:
        raise ValueError("Design rules changed since batch execution")
    template = ROOT / "templates" / rules["expression"]["backbone_template"]
    if digest(template) != manifest["template_sha256"]:
        raise ValueError("Backbone template changed since batch execution")

    checks = []
    order_usage: dict[str, set[int]] = defaultdict(set)
    order_suffixes: dict[str, str] = {}
    for index, item in enumerate(provenance["items"]):
        construct_id = numbering.construct_id(index)
        if item.get("construct_id") != construct_id or item.get("map") != f"maps/{plasmid_prefix}-{construct_id}.dna":
            raise ValueError(f"Provenance numbering mismatch: {item['candidate_id']}")
        source = Path(item["source"])
        structure = project / item["structure"]
        if digest(source) != item["sha256"] or digest(structure) != item["sha256"]:
            raise ValueError(f"Input source/copy mismatch: {item['candidate_id']}")
        design = read_structure_design(structure, construct_id, protein_chain=manifest["parameters"]["chain"])
        map_path = project / "maps" / f"{plasmid_prefix}-{construct_id}.dna"
        record = SeqIO.read(map_path, "snapgene")
        poi = get_poi_feature(record)
        cds = str(poi.extract(record.seq)).upper()
        codons = [cds[start : start + 3] for start in range(0, len(cds), 3)]
        translated = "".join(
            "X" if codon == rules["input"]["ncaa_codon"] else str(Seq(codon).translate())
            for codon in codons
        )
        if translated != design.expression_sequence:
            raise ValueError(f"POI translation mismatch: {construct_id}")
        observed_ncaa = [
            position for position, codon in enumerate(codons, 1)
            if codon == rules["input"]["ncaa_codon"]
        ]
        if observed_ncaa != list(design.ncaa_expression_positions):
            raise ValueError(f"ncAA codon position mismatch: {construct_id}")
        tail = str(record.seq[int(poi.location.end):int(poi.location.end) + len(tail_dna)]).upper()
        if tail != tail_dna or str(Seq(tail).translate()) != expected_tail:
            raise ValueError(f"Fixed C-terminal cassette mismatch: {construct_id}")
        if design.expression_sequence.count("C") != item["native_cysteines"]:
            raise ValueError(f"POI cysteine count changed: {construct_id}")
        primer_count = 0
        for feature in record.features:
            if feature.type != "primer_bind":
                continue
            primer_count += 1
            primer = str(feature.extract(record.seq)).upper()
            label = feature.qualifiers.get("label", [""])[0]
            order_usage[primer].add(index + 1)
            if label.startswith("vec-"):
                suffix = label
            elif label.startswith(f"{construct_id}-"):
                suffix = label[len(construct_id) + 1:]
            else:
                raise ValueError(f"Primer label does not match construct: {construct_id}: {label}")
            if primer in order_suffixes and order_suffixes[primer] != suffix:
                raise ValueError(f"Reused primer has inconsistent roles: {label}")
            order_suffixes[primer] = suffix
            if len(primer) < 15:
                raise ValueError(f"Primer shorter than 15 nt: {construct_id}")
            motif = primer[-15:]
            target = motif if feature.location.strand == 1 else str(Seq(motif).reverse_complement())
            if circular_count(str(record.seq).upper(), target) != 1:
                raise ValueError(f"Nonunique primer 3' 15 nt: {construct_id}: {feature.qualifiers.get('label')}")
        checks.append({
            "candidate_id": item["candidate_id"], "construct_id": construct_id,
            "map": str(map_path.relative_to(project)), "plasmid_bp": len(record.seq),
            "protein_aa": len(design.expression_sequence),
            "ncaa_expression_positions": observed_ncaa, "primer_features": primer_count,
            "codon_strategy": manifest["selected_codon_strategies"][construct_id],
        })

    workbook = load_workbook(project / "all_primers.xlsx", read_only=True)
    primers = list(workbook.active.values)[1:]
    workbook.close()
    if len({row[1] for row in primers}) != len(primers):
        raise ValueError("Ordering workbook contains duplicate oligonucleotide sequences")
    order_by_sequence = {row[1]: row[0] for row in primers}
    if set(order_by_sequence) != set(order_usage):
        raise ValueError("Ordering workbook does not cover exactly the map primer sequences")
    if rules["oligos"].get("order_names_include_construct_usage", False):
        all_numbers = set(range(1, len(checks) + 1))
        for sequence, name in order_by_sequence.items():
            numbers = sorted(order_usage[sequence])
            group = (
                f"{numbers[0]}to{numbers[-1]}"
                if set(numbers) == all_numbers and len(numbers) >= 5
                else "_".join(map(str, numbers))
            )
            expected_name = f"{group}-{order_suffixes[sequence]}"
            if name != expected_name:
                raise ValueError(f"Ordering name mismatch: {name} != {expected_name}")
    summary = load_workbook(project / "maps" / "statistics_summary.xlsx", read_only=True)
    rows = list(summary.active.values)[1:]
    summary.close()
    if len(rows) != len(checks) or any(row[21] != "PASS" for row in rows):
        raise ValueError("Statistics workbook does not contain one PASS row per design")
    return {
        "pass": True, "designs": len(checks), "unique_orderable_oligos": len(primers),
        "maps_checked": checks,
        "checks": "Source/output SHA-256, PDB sequence, POI/TAG/tail translation, primer 3' 15-nt uniqueness, exact order sequence coverage, shared-use order names and deduplication",
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, required=True)
    args = parser.parse_args()
    project = args.project if args.project.is_absolute() else ROOT / args.project
    result = validate(project)
    output = project / "batch_qc.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(f"PASS: {result['designs']} designs, {result['unique_orderable_oligos']} unique oligos; {output}")
